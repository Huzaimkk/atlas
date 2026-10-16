import openpyxl, json, sys

wb = openpyxl.load_workbook('Opportunity_Research_EU_USA_UAE.xlsx', data_only=True)

def sheet_rows(name, header_row):
    ws = wb[name]
    headers = [ws.cell(row=header_row, column=c).value for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        row = {}
        empty = True
        for c, h in enumerate(headers, 1):
            if h is None:
                continue
            v = ws.cell(row=r, column=c).value
            if v is not None:
                empty = False
            row[h] = v
        if not empty:
            rows.append(row)
    return rows

def default_str(v):
    if hasattr(v, 'isoformat'):
        return v.isoformat()
    return v

opportunities = sheet_rows('Opportunities', 5)
corrections = sheet_rows('Corrections', 5)
texas_focus = sheet_rows('Texas focus', 5)
university_details = sheet_rows('University details', 5)
employer_targets = sheet_rows('Employer targets', 5)
events = sheet_rows('Events', 5)

out = {
    'opportunities': opportunities,
    'corrections': corrections,
    'texas_focus': texas_focus,
    'university_details': university_details,
    'employer_targets': employer_targets,
    'events': events,
}

with open(sys.argv[1], 'w', encoding='utf-8') as f:
    json.dump(out, f, default=default_str, ensure_ascii=False, indent=1)

print(f"Opportunities: {len(opportunities)}")
print(f"Corrections: {len(corrections)}")
print(f"Texas focus: {len(texas_focus)}")
print(f"University details: {len(university_details)}")
print(f"Employer targets: {len(employer_targets)}")
print(f"Events: {len(events)}")
